from __future__ import annotations

from io import BytesIO

from docx import Document


def generate_word_report(
    *,
    company_name: str,
    assessment_name: str,
    framework_name: str,
    responses,
    domain_scores,
    executive_summary=None,
    recommendations=None,
):
    doc = Document()

    doc.add_heading("GRC Assessment Report", level=0)
    doc.add_paragraph(f"Company: {company_name}")
    doc.add_paragraph(f"Assessment: {assessment_name}")
    doc.add_paragraph(f"Framework: {framework_name}")

    doc.add_heading("Domain Scores", level=1)
    if domain_scores:
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Domain"
        hdr[1].text = "Score"

        for domain_name, score in domain_scores.items():
            row = table.add_row().cells
            row[0].text = str(domain_name)
            row[1].text = str(score)
    else:
        doc.add_paragraph("No domain scores available.")

    doc.add_heading("Executive Summary", level=1)
    if executive_summary:
        if getattr(executive_summary, "summary_text", ""):
            doc.add_paragraph(getattr(executive_summary, "summary_text"))
        if getattr(executive_summary, "strengths_text", ""):
            doc.add_heading("Strengths", level=2)
            doc.add_paragraph(getattr(executive_summary, "strengths_text"))
        if getattr(executive_summary, "gaps_text", ""):
            doc.add_heading("Gaps", level=2)
            doc.add_paragraph(getattr(executive_summary, "gaps_text"))
        if getattr(executive_summary, "recommendations_text", ""):
            doc.add_heading("Management Recommendations", level=2)
            doc.add_paragraph(getattr(executive_summary, "recommendations_text"))
    else:
        doc.add_paragraph("No executive summary available.")

    doc.add_heading("Recommendations", level=1)
    if recommendations:
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr = table.rows[0].cells
        hdr[0].text = "Priority"
        hdr[1].text = "Title"
        hdr[2].text = "Description"
        hdr[3].text = "Domain"

        for rec in recommendations:
            row = table.add_row().cells
            row[0].text = str(getattr(rec, "priority", "") or "")
            row[1].text = str(getattr(rec, "title", "") or "")
            row[2].text = str(getattr(rec, "description", "") or "")
            row[3].text = str(getattr(rec, "domain_name", "") or "")
    else:
        doc.add_paragraph("No recommendations available.")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()