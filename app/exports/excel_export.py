from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def generate_excel_report(
    *,
    company_name: str,
    assessment_name: str,
    framework_name: str,
    responses,
    domain_scores,
    executive_summary=None,
    recommendations=None,
):
    wb = Workbook()

    ws_cover = wb.active
    ws_cover.title = "Summary"

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    bold_font = Font(bold=True)

    ws_cover["A1"] = "GRC Assessment Report"
    ws_cover["A1"].font = Font(bold=True, size=14)

    ws_cover["A3"] = "Company"
    ws_cover["B3"] = company_name
    ws_cover["A4"] = "Assessment"
    ws_cover["B4"] = assessment_name
    ws_cover["A5"] = "Framework"
    ws_cover["B5"] = framework_name

    ws_cover["A7"] = "Executive Summary"
    ws_cover["A7"].font = bold_font
    ws_cover["A8"] = getattr(executive_summary, "summary_text", "") if executive_summary else ""

    ws_cover["A10"] = "Strengths"
    ws_cover["A10"].font = bold_font
    ws_cover["A11"] = getattr(executive_summary, "strengths_text", "") if executive_summary else ""

    ws_cover["A13"] = "Gaps"
    ws_cover["A13"].font = bold_font
    ws_cover["A14"] = getattr(executive_summary, "gaps_text", "") if executive_summary else ""

    ws_cover["A16"] = "Recommendations"
    ws_cover["A16"].font = bold_font
    ws_cover["A17"] = getattr(executive_summary, "recommendations_text", "") if executive_summary else ""

    ws_cover.column_dimensions["A"].width = 22
    ws_cover.column_dimensions["B"].width = 80

    ws_scores = wb.create_sheet("Domain Scores")
    ws_scores.append(["Domain", "Score"])
    for cell in ws_scores[1]:
        cell.fill = header_fill
        cell.font = header_font

    for domain_name, score in (domain_scores or {}).items():
        ws_scores.append([domain_name, score])

    ws_scores.column_dimensions["A"].width = 40
    ws_scores.column_dimensions["B"].width = 12

    ws_answers = wb.create_sheet("Answers")
    ws_answers.append(
        [
            "Domain",
            "Question ID",
            "Question",
            "Answer",
            "Score",
            "Weight",
            "Risk",
            "Notes",
            "Proof",
        ]
    )
    for cell in ws_answers[1]:
        cell.fill = header_fill
        cell.font = header_font

    for row in responses or []:
        ws_answers.append(
            [
                row.get("domain_name") or row.get("domain"),
                row.get("question_id"),
                row.get("question_text") or row.get("question"),
                row.get("answer_value") or row.get("selected_value"),
                row.get("score"),
                row.get("weight"),
                row.get("risk"),
                row.get("notes") or row.get("comment"),
                row.get("proof") or row.get("evidence"),
            ]
        )

    ws_answers.column_dimensions["A"].width = 28
    ws_answers.column_dimensions["B"].width = 16
    ws_answers.column_dimensions["C"].width = 60
    ws_answers.column_dimensions["D"].width = 18
    ws_answers.column_dimensions["E"].width = 10
    ws_answers.column_dimensions["F"].width = 10
    ws_answers.column_dimensions["G"].width = 12
    ws_answers.column_dimensions["H"].width = 40
    ws_answers.column_dimensions["I"].width = 30

    ws_recs = wb.create_sheet("Recommendations")
    ws_recs.append(["Priority", "Title", "Description", "Domain", "Score", "Source"])
    for cell in ws_recs[1]:
        cell.fill = header_fill
        cell.font = header_font

    for rec in recommendations or []:
        ws_recs.append(
            [
                getattr(rec, "priority", ""),
                getattr(rec, "title", ""),
                getattr(rec, "description", ""),
                getattr(rec, "domain_name", ""),
                getattr(rec, "score", None),
                getattr(rec, "source", ""),
            ]
        )

    ws_recs.column_dimensions["A"].width = 14
    ws_recs.column_dimensions["B"].width = 35
    ws_recs.column_dimensions["C"].width = 70
    ws_recs.column_dimensions["D"].width = 28
    ws_recs.column_dimensions["E"].width = 10
    ws_recs.column_dimensions["F"].width = 12

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()
